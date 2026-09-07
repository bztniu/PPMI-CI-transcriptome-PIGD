import pandas as pd, openpyxl
from openpyxl.styles import Font
ROOT = r"E:/PPMI帕金森数据库专用"
OUT  = ROOT + r"\结果\新结果\实验\movementdisorders\new\revision_v2\2026-08-31-第四版-评审修订"
xp   = OUT + r"\Supplementary_Tables_updated.xlsx"
wb = openpyxl.load_workbook(xp)
print("before:", wb.sheetnames[-4:])

# --- TableS25: CBC validation ---
cbc = pd.read_csv(OUT + r"\cbc_vs_deconvolution.csv")
ifr = pd.read_csv(ROOT + r"\结果\新结果\实验\movementdisorders\new\revision_v2\figures_pc1\data\immune_fractions.csv")
cbc = cbc.merge(ifr[["SAMPLE_ID","group"]], left_on="PATNO",
                right_on=ifr["SAMPLE_ID"].map(lambda s: s).index if False else ifr["SAMPLE_ID"], how="left") if False else cbc
pmap = pd.read_csv(ROOT + r"\结果\新结果\实验\movementdisorders\new\revision_v2\data\PD_all_clustering_methods.csv")
sid2pat = dict(zip(pmap["SAMPLE_ID"], pmap["PATNO"]))
ifr["PATNO"] = ifr["SAMPLE_ID"].map(sid2pat)
cbc = cbc.merge(ifr[["PATNO","group"]].drop_duplicates("PATNO"), on="PATNO", how="left")
if "TableS25" not in wb.sheetnames:
    ws = wb.create_sheet("TableS25")
    ws.append(["Table S25. Validation of NNLS-deconvolution immune fractions against measured complete blood count differentials (PPMI, screening visit; n = 386). Estimated lymphocyte fraction is the sum of all lymphoid cell-type estimates."])
    ws["A1"].font = Font(bold=True)
    ws.append(["PATNO","CI group","Measured neutrophils (%)","Measured lymphocytes (%)","Measured monocytes (%)","Estimated neutrophil fraction","Estimated lymphoid fraction","Estimated monocyte fraction"])
    for _, r in cbc.iterrows():
        ws.append([int(r["PATNO"]), r.get("group"), round(r["Neutrophils (%)"],1), round(r["Lymphocytes (%)"],1), round(r["Monocytes (%)"],1), round(r["Neut"],4), round(r["Lymph"],4), round(r["Monoc"],4)])
    print("TableS25 rows:", ws.max_row)

# --- TableS26: null test ---
nul = pd.read_csv(OUT + r"\random66_null_results.csv")
if "TableS26" not in wb.sheetnames:
    ws = wb.create_sheet("TableS26")
    ws.append(["Table S26. Gene-set null test: year-5 PIGD interaction (continuous score specification) for 1,000 random 66-gene sets drawn from the 28,560 expressed non-scoring genes. Observed CI-score estimate: beta = -0.183; empirical p = 0.007 (6 of 1,000 random sets reached the observed absolute estimate)."])
    ws["A1"].font = Font(bold=True)
    ws.append(["Random set index","Year-5 interaction beta","Nominal p"])
    for _, r in nul.iterrows():
        ws.append([int(r["iter"]), round(r["beta"],4), round(r["p"],5)])
    print("TableS26 rows:", ws.max_row)

# --- INDEX ---
idx = wb["INDEX"]
existing = [row[0] for row in idx.iter_rows(values_only=True)]
if not any("S25" in str(v) for v in existing if v):
    idx.append(["TableS25","Validation of deconvolution-derived immune fractions against measured CBC differentials (n = 386)"])
if not any("S26" in str(v) for v in existing if v):
    idx.append(["TableS26","Gene-set null test for the CI-PIGD year-5 interaction (1,000 random 66-gene sets)"])
print("INDEX rows:", idx.max_row)
wb.save(xp)
print("saved")
